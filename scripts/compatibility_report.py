"""Record local lock/artifact evidence. This is compatibility, never a quality benchmark."""
import hashlib
import importlib.metadata as metadata
import io
import json
import platform
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
writer = PdfWriter()
writer.add_blank_page(width=72, height=72)
pdf = io.BytesIO()
writer.write(pdf)
assert len(PdfReader(io.BytesIO(pdf.getvalue())).pages) == 1
source = Document()
source.add_paragraph('EvidenceDesk parser compatibility.')
docx = io.BytesIO()
source.save(docx)
assert Document(io.BytesIO(docx.getvalue())).paragraphs[0].text == 'EvidenceDesk parser compatibility.'
assert BeautifulSoup('<p>Evidence</p>', 'html.parser').get_text() == 'Evidence'
licenses = []
for dist in metadata.distributions():
    raw = dist.metadata.get('License-Expression') or dist.metadata.get('License') or ''
    classifiers = [s for s in dist.metadata.get_all('Classifier', []) if s.startswith('License ::')]
    licenses.append({'name': dist.metadata['Name'], 'version': dist.version, 'declared_license': raw[:2000], 'classifiers': classifiers})
report = {
    'checked_on': '2026-10-11', 'platform': platform.platform(), 'python': platform.python_version(),
    'checks': {'pdf_read_write': 'pass', 'docx_text_read_write': 'pass', 'html_parser': 'pass'},
    'locks_sha256': {p: hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in ['uv.lock', 'pnpm-lock.yaml', 'infra/images.lock', 'contracts/config/model-artifacts.lock.json']},
    'python_dependency_licenses': sorted(licenses, key=lambda d: d['name'].lower()),
    'advisories': {
        'python': json.loads(Path('/tmp/cbot-pip-audit-full.json').read_text()),
        'javascript': json.loads(Path('/tmp/cbot-pnpm-audit.json').read_text()),
    },
    'scope': 'Local Linux x86_64 runtime, parser and model compatibility; vendor probes recorded separately. No live generative model quality claim.'
}
(ROOT/'docs/10-testing/build-compatibility-results.json').write_text(json.dumps(report, indent=2)+'\n')
print('Compatibility evidence written; packages:', len(licenses))
