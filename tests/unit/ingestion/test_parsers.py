import hashlib
import json

import pytest
from docx import Document
from evidencedesk.errors import Problem
from evidencedesk.ingestion.parsers.core import extract
from evidencedesk.ingestion.validation import MIMES
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def parse(path):
    return extract(
        path, path.name, MIMES[path.suffix], hashlib.sha256(path.read_bytes()).hexdigest()
    )


def make_pdf(path, texts, *, encrypted=False):
    writer = PdfWriter()
    for value in texts:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 30 250 Td ({value}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("fixture-password")
    writer.write(path)


@pytest.fixture
def golden_sources(tmp_path):
    # Expectations are hand-authored; output is never used to regenerate expected text.
    fixtures = [
        (
            "plain.txt",
            "First paragraph.\n\nSecond paragraph.",
            ["First paragraph.", "Second paragraph."],
        ),
        ("unicode.txt", "Cafe\u0301\r\nsecond line\t \n", ["Café\nsecond line"]),
        (
            "headings.md",
            "# Rules\n\nAlways cite evidence.\n\n## Limits\n\nStay within budget.",
            ["# Rules", "Always cite evidence.", "## Limits", "Stay within budget."],
        ),
        (
            "code.md",
            '# Code\n\n```python\n  print("safe")\n```\n',
            ["# Code", '```python\n  print("safe")\n```'],
        ),
        (
            "safe.html",
            '<h1>Rules</h1><script>secret()</script><style>bad</style><p>Hello <b>world</b>.</p><img src="http://example.invalid/tracker"><!-- hidden -->',
            ["Rules", "Hello world."],
        ),
        (
            "table.csv",
            'Name,Rule\nA,"Cite, then explain"\nB,Keep private\n',
            ["Name", "Rule", "A", "Cite, then explain", "B", "Keep private"],
        ),
        (
            "values.json",
            '{"a/b":["one",2,true,null],"unicode":"Café"}',
            ["one", "2", "true", "null", "Café"],
        ),
    ]
    result = []
    for name, text, expected in fixtures:
        path = tmp_path / name
        path.write_bytes(text.encode())
        result.append((path, expected))
    path = tmp_path / "text.pdf"
    make_pdf(
        path,
        ["First page has enough visible source text.", "Second page has its own source locator."],
    )
    result.append(
        (
            path,
            [
                "First page has enough visible source text.",
                "Second page has its own source locator.",
            ],
        )
    )
    path = tmp_path / "paragraphs.docx"
    document = Document()
    document.add_heading("Policy", 1)
    document.add_paragraph("Keep evidence private.")
    document.save(path)
    result.append((path, ["Policy", "Keep evidence private."]))
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip(
        [c for r in table.rows for c in r.cells], ["Name", "Rule", "A", "Cite sources"], strict=True
    ):
        cell.text = value
    document.save(path)
    result.append((path, ["Name", "Rule", "A", "Cite sources"]))
    return result


def test_ten_golden_sources_exact_text_offsets_and_provenance(golden_sources):
    for path, expected in golden_sources:
        result = parse(path)
        assert [block.text for block in result.blocks] == expected, path.name
        assert parse(path).model_dump() == result.model_dump()
        stream = "\n\n".join(expected)
        for block in result.blocks:
            assert stream[block.char_start : block.char_end] == block.text
            assert (
                block.locator.page is None
                if path.suffix != ".pdf"
                else block.locator.page in (1, 2)
            )
        if path.suffix == ".pdf":
            assert result.page_count == 2 and [b.locator.page for b in result.blocks] == [1, 2]
        if path.name == "unicode.txt":
            assert result.blocks[0].original_text.startswith("Cafe\u0301")
        if path.name == "headings.md":
            assert result.blocks[-1].heading_path == ["Rules", "Limits"]
            assert result.blocks[-1].locator.line_start == 7
        if path.name == "table.csv":
            assert result.blocks[-1].header == "Rule"
            assert (result.blocks[-1].locator.row, result.blocks[-1].locator.column) == (3, 2)
        if path.name == "values.json":
            assert result.blocks[0].locator.pointer == "/a~1b/0"
        if path.suffix == ".html":
            assert result.omitted_images == 1 and result.warnings


@pytest.mark.parametrize(
    "texts,encrypted,code",
    [
        (["", ""], False, "ocr_required"),
        (["Enough extractable text for a real page.", ""], False, "partial_coverage"),
        (["Enough extractable text for a real page."], True, "encrypted_pdf"),
        (["text"] * 201, False, "page_limit"),
    ],
)
def test_pdf_rejects_incomplete_encrypted_or_oversized(tmp_path, texts, encrypted, code):
    path = tmp_path / "reject.pdf"
    make_pdf(path, texts, encrypted=encrypted)
    with pytest.raises(Problem) as exc:
        parse(path)
    assert exc.value.code == code


@pytest.mark.parametrize(
    "value,code",
    [
        ("[" * 21 + "0" + "]" * 21, "json_depth"),
        ('{"a":1,"a":2}', "duplicate_json_key"),
        ('{"a":NaN}', "invalid_json"),
        ("{bad", "parse_failed"),
    ],
)
def test_json_structure_limits(tmp_path, value, code):
    path = tmp_path / "input.json"
    path.write_text(value)
    with pytest.raises(Problem) as exc:
        parse(path)
    assert exc.value.code == code


def test_wide_csv_retains_cell_locators_and_limits(tmp_path):
    path = tmp_path / "wide.csv"
    path.write_text(",".join("cell" for _ in range(1000)))
    assert len(parse(path).blocks) == 1000
    path.write_text(",".join("cell" for _ in range(1001)))
    with pytest.raises(Problem) as exc:
        parse(path)
    assert exc.value.code == "table_limit"


def test_malformed_pdf_and_json_braces_in_strings(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF-1.7 not a valid document")
    with pytest.raises(Problem) as exc:
        parse(path)
    assert exc.value.code == "parse_failed"
    path = tmp_path / "braces.json"
    path.write_text(json.dumps({"value": "[" * 30}))
    assert parse(path).blocks[0].text == "[" * 30


def test_html_table_coordinates_and_json_decimal_fidelity(tmp_path):
    path = tmp_path / "table.html"
    path.write_text(
        "<table><tr><th>Name</th><th>Rule</th></tr><tr><td>A</td><td>Cite <b>sources</b>.</td></tr></table>"
    )
    result = parse(path)
    assert [b.text for b in result.blocks] == ["Name", "Rule", "A", "Cite sources."]
    assert result.blocks[-1].header == "Rule"
    assert (result.blocks[-1].locator.row, result.blocks[-1].locator.column) == (2, 2)
    path = tmp_path / "decimal.json"
    path.write_text('{"exact": 0.123456789012345678901234567890}')
    assert parse(path).blocks[0].text == "0.123456789012345678901234567890"
